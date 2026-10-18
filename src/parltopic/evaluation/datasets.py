"""Build frozen, labelled evaluation sets from the OpenParlData hackathon files.

E1 digitalisation: PolitMonitor positives vs. untagged affairs of the same parliaments,
    period and affair types (dev 2018-2023 for tuning, test 2024-2026 frozen).
E2 Curia Vista subject areas: multi-label federal affairs, test 2024-2026.
E4 DVS: the 30 affairs listed by Digital Public Services Switzerland (positives only).
"""

import json
import random
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from parltopic.ingestion.opendata_export import (
    available_languages,
    load_bodies,
    load_che_texts,
    read_csv,
    read_ndjson,
    to_affair,
)

SPLITS = {"dev": ("2018-01-01", "2023-12-31"), "test": ("2024-01-01", "2026-12-31")}
MIN_TYPE_POSITIVES = 10


def _in_split(date: str | None, split: str) -> bool:
    start, end = SPLITS[split]
    return bool(date) and start <= date[:10] <= end


def _che_language(record: dict[str, Any], rng: random.Random) -> str | None:
    languages = [lang for lang in available_languages(record) if lang in ("de", "fr", "it")]
    return rng.choice(languages) if languages else None


def _rows(
    records: list[dict[str, Any]],
    bodies: dict[str, dict[str, Any]],
    texts: dict[int, dict[str, str]],
    rng: random.Random,
    extra: dict[int, dict[str, Any]],
) -> list[dict[str, Any]]:
    rows = []
    for record in records:
        body = bodies.get(record["body_key"])
        language = _che_language(record, rng) if record["body_key"] == "CHE" else None
        text = texts.get(record["id"], {}).get(language or "", "")
        affair = to_affair(record, body, language=language, text=text)
        if affair:
            rows.append({**affair.model_dump(mode="json"), **extra[record["id"]]})
    return rows


def _write(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as file:
        for row in rows:
            file.write(json.dumps(row, ensure_ascii=False) + "\n")


def build_eval_sets(
    raw_dir: str | Path,
    out_dir: str | Path,
    seed: int = 42,
    sizes: dict[str, tuple[int, int]] | None = None,
    curia_size: int = 500,
    eligibility: str = "historical",
) -> dict[str, Any]:
    raw, out = Path(raw_dir), Path(out_dir)
    rng = random.Random(seed)
    sizes = sizes or {"dev": (150, 300), "test": (400, 800)}
    bodies = load_bodies(raw)
    positive_ids = {
        int(r["openparldata_id"]) for r in read_csv(raw / "politmonitor_digitalisierung.csv")
    }
    dvs_ids = {int(r["openparldata_id"]) for r in read_csv(raw / "dvs_digitalisierung.csv")}
    curia = {
        row["openparldata_id"]: row
        for row in read_ndjson(raw / "che_affairs_keywords.jsonl")
        if row["topic_ids"]
    }

    affairs = {record["id"]: record for record in read_ndjson(raw / "affairs.ndjson.gz")}
    positives = [affairs[i] for i in positive_ids if i in affairs]
    # "historical" reproduces the frozen files in data/eval byte for byte: eligible parliaments
    # and affair types come from all positives. "dev" learns them from the development period
    # only, so later positives cannot leak into the sampler; use it for new sealed holdouts.
    if eligibility == "dev":
        source = [r for r in positives if _in_split(r.get("begin_date"), "dev")]
    elif eligibility == "historical":
        source = positives
    else:
        raise ValueError(f"eligibility must be 'historical' or 'dev', not {eligibility!r}")
    positive_bodies = Counter(record["body_key"] for record in source)
    type_counts = Counter(record["type_harmonized_en"] for record in source)
    positive_types = {t for t, n in type_counts.items() if n >= MIN_TYPE_POSITIVES}

    meta: dict[str, Any] = {
        "seed": seed,
        "positive_types": sorted(map(str, positive_types)),
        "eligibility_source": (
            "dev positives only" if eligibility == "dev" else "all positives (historical)"
        ),
        "label_semantics": "positive_unlabeled_proxy",
        "benchmark_status": "opened_validation_not_sealed",
        "text_source": (
            "texts_CHE.ndjson.gz" if (raw / "texts_CHE.ndjson.gz").exists()
            else "che_affair_texts.ndjson.gz"
        ),
        "text_order": "submitted text, reasoning, lead; stable text ID within type",
    }
    chosen: dict[str, dict[str, list[dict[str, Any]]]] = {}
    for split in SPLITS:
        eligible = [
            record
            for record in affairs.values()
            if record["body_key"] in positive_bodies
            and record["type_harmonized_en"] in positive_types
            and _in_split(record.get("begin_date"), split)
            and available_languages(record)
        ]
        split_pos = [r for r in eligible if r["id"] in positive_ids]
        split_neg = [r for r in eligible if r["id"] not in positive_ids]
        n_pos, n_neg = sizes[split]
        sample_pos = rng.sample(split_pos, min(n_pos, len(split_pos)))
        # Negatives per parliament in proportion to its positives, so the parliament alone
        # does not reveal the label.
        neg_by_body = defaultdict(list)
        for record in split_neg:
            neg_by_body[record["body_key"]].append(record)
        pos_by_body = Counter(r["body_key"] for r in sample_pos)
        sample_neg = []
        for body_key, count in pos_by_body.items():
            quota = round(n_neg * count / len(sample_pos))
            pool = neg_by_body[body_key]
            sample_neg.extend(rng.sample(pool, min(quota, len(pool))))
        chosen[split] = {"pos": sample_pos, "neg": sample_neg}
        if split == "dev":
            dev_pool = (split_pos, split_neg)
        meta[f"e1_{split}"] = {
            "eligible": len(eligible),
            "eligible_positives": len(split_pos),
            "prevalence": round(len(split_pos) / len(eligible), 4) if eligible else None,
            "positives": len(sample_pos),
            "negatives": len(sample_neg),
        }

    curia_test = [
        affairs[i]
        for i in curia
        if i in affairs and _in_split(affairs[i].get("begin_date"), "test")
    ]
    curia_sample = rng.sample(curia_test, min(curia_size, len(curia_test)))
    dvs = [affairs[i] for i in dvs_ids if i in affairs]

    che_ids = {
        r["id"]
        for group in [*(c["pos"] + c["neg"] for c in chosen.values()), curia_sample, dvs]
        for r in group
        if r["body_key"] == "CHE"
    }
    texts = load_che_texts(raw, che_ids)

    for split, groups in chosen.items():
        extra = {r["id"]: {"label": True, "split": split} for r in groups["pos"]}
        extra |= {r["id"]: {"label": False, "split": split} for r in groups["neg"]}
        rows = _rows(groups["pos"] + groups["neg"], bodies, texts, rng, extra)
        rng.shuffle(rows)
        _write(out / f"e1_digitalisation_{split}.jsonl", rows)
        meta[f"e1_{split}"]["languages"] = dict(Counter(r["language"] for r in rows))
        meta[f"e1_{split}"]["levels"] = dict(Counter(r["level"] for r in rows))

    curia_extra = {
        i: {
            "labels": [str(t) for t in curia[i]["topic_ids"]],
            "label_names": curia[i]["topics_de"],
            "split": "test",
        }
        for i in (r["id"] for r in curia_sample)
    }
    curia_rows = _rows(curia_sample, bodies, texts, rng, curia_extra)
    _write(out / "e2_curia_test.jsonl", curia_rows)
    meta["e2_test"] = {
        "affairs": len(curia_rows),
        "languages": dict(Counter(r["language"] for r in curia_rows)),
        "labels_per_affair": round(sum(len(r["labels"]) for r in curia_rows) / len(curia_rows), 2),
    }

    dvs_rows = _rows(
        dvs, bodies, texts, rng, {r["id"]: {"label": True, "split": "dvs"} for r in dvs}
    )
    _write(out / "e4_dvs.jsonl", dvs_rows)
    meta["e4_dvs"] = {
        "affairs": len(dvs_rows),
        "levels": dict(Counter(r["level"] for r in dvs_rows)),
    }

    # Few-shot pool: labelled affairs from the dev period that are NOT in the dev sample, so a
    # dev or test affair never sees its own label. Own generator: the sets above stay identical.
    pool_rng = random.Random(seed + 1)
    dev_ids = {r["id"] for r in chosen["dev"]["pos"] + chosen["dev"]["neg"]}
    pool_pos = [r for r in dev_pool[0] if r["id"] not in dev_ids]
    pool_neg = pool_rng.sample(
        [r for r in dev_pool[1] if r["id"] not in dev_ids], min(4 * len(pool_pos), len(dev_pool[1]))
    )
    pool_extra = {r["id"]: {"label": True, "split": "pool"} for r in pool_pos}
    pool_extra |= {r["id"]: {"label": False, "split": "pool"} for r in pool_neg}
    pool_rows = _rows(pool_pos + pool_neg, bodies, {}, pool_rng, pool_extra)
    _write(out / "e1_fewshot_pool.jsonl", pool_rows)
    meta["e1_fewshot_pool"] = {"positives": len(pool_pos), "negatives": len(pool_neg)}

    (out / "meta.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")
    return meta


def export_aps_scheme(raw_dir: str | Path, out_path: str | Path) -> list[dict[str, Any]]:
    """APS scheme (domain > chapter > topic, DE/FR) from aps_topics.xlsx as JSON."""
    from openpyxl import load_workbook

    sheet = load_workbook(Path(raw_dir) / "aps_topics.xlsx", read_only=True)["Topics"]
    header, *rows = list(sheet.iter_rows(values_only=True))
    topics = [dict(zip(header, row, strict=True)) for row in rows if row and row[0]]
    for topic in topics:
        topic["chapter_code"] = ".".join(str(topic["code"]).split(".")[:2])
    path = Path(out_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(topics, indent=1, ensure_ascii=False), encoding="utf-8")
    return topics


def build_aps_sample(
    raw_dir: str | Path,
    eval_dir: str | Path,
    seed: int = 7,
    federal: int = 25,
    cantonal: dict[str, int] | None = None,
) -> list[dict[str, Any]]:
    """E3: 50 affairs from 2024-2026 for hand labelling with APS chapter > topic.

    Federal affairs come from the E2 Curia sample (DE/FR/IT mixed, Curia areas kept for a
    later cross-check, not shown to labellers); cantonal ones are drawn per language.
    """
    raw, out = Path(raw_dir), Path(eval_dir)
    rng = random.Random(seed)
    cantonal = cantonal or {"de": 13, "fr": 8, "it": 4}
    curia_rows = [row for row in _read_jsonl(out / "e2_curia_test.jsonl") if row["title"]]
    rows = [{**r, "split": "aps"} for r in rng.sample(curia_rows, federal)]
    bodies = load_bodies(raw)
    seen = {r["affair_id"] for r in rows}
    pools: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for record in read_ndjson(raw / "affairs.ndjson.gz"):
        body = bodies.get(record["body_key"], {})
        if body.get("type") != "canton" or not _in_split(record.get("begin_date"), "test"):
            continue
        if record.get("type_harmonized_en") not in {"Motion", "Postulate", "Interpellation"}:
            continue
        affair = to_affair(record, body)
        if affair and len(affair.title) >= 25 and affair.affair_id not in seen:
            pools[affair.language].append(affair.model_dump(mode="json"))
    for language, count in cantonal.items():
        rows.extend({**r, "split": "aps"} for r in rng.sample(pools[language], count))
    rng.shuffle(rows)
    _write(out / "e3_aps_sample.jsonl", rows)
    return rows


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    with path.open(encoding="utf-8") as file:
        return [json.loads(line) for line in file if line.strip()]
